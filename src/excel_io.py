import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string

from src.config import Config
from src.validacion import Decision


@dataclass(frozen=True)
class RegistroExcel:
    fila: int
    nit: str
    nombre: str
    valor: Decimal


@dataclass(frozen=True)
class FilaInvalida:
    fila: int
    motivo: str


def parsear_nit(valor) -> str:
    """3989120 / '3.989.120' -> '3989120'. Rechaza letras y guiones (p. ej. con DV)."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if not re.fullmatch(r"[\d.,\s]+", texto):
        raise ValueError(f"NIT no válido: {texto!r}")
    return "".join(c for c in texto if c.isdigit())


def parsear_valor_excel(valor) -> Decimal:
    """Convierte el valor de la celda a Decimal. Rechaza formatos ambiguos."""
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    texto = str(valor).strip().replace(" ", "")
    if not texto:
        raise ValueError("valor vacío")
    if "," in texto and "." in texto:
        decimal = "," if texto.rfind(",") > texto.rfind(".") else "."
        miles = "." if decimal == "," else ","
        texto = texto.replace(miles, "").replace(decimal, ".")
    elif texto.count(",") > 1 or texto.count(".") > 1:
        texto = texto.replace(",", "").replace(".", "")
    elif "," in texto or "." in texto:
        raise ValueError(f"formato ambiguo: {texto!r}")
    try:
        return Decimal(texto)
    except InvalidOperation:
        raise ValueError(f"valor no numérico: {texto!r}")


def leer_registros(config: Config) -> tuple[list[RegistroExcel], list[FilaInvalida]]:
    """Lee el Excel de entrada. Las filas que no se entienden se devuelven aparte."""
    libro = load_workbook(config.archivo_excel, data_only=True)
    hoja = libro[config.hoja_excel] if config.hoja_excel else libro.active
    indices = [
        column_index_from_string(c)
        for c in (config.col_nit, config.col_nombre, config.col_valor)
    ]

    registros: list[RegistroExcel] = []
    invalidas: list[FilaInvalida] = []
    filas = hoja.iter_rows(min_row=config.fila_inicio, values_only=True)
    for numero, fila in enumerate(filas, start=config.fila_inicio):
        celdas = [fila[i - 1] if i - 1 < len(fila) else None for i in indices]
        if all(c in (None, "") for c in celdas):
            continue  # fila vacía
        nit_raw, nombre_raw, valor_raw = celdas
        if nit_raw in (None, "") and nombre_raw in (None, ""):
            invalidas.append(
                FilaInvalida(fila=numero, motivo="sin NIT ni nombre (¿fila de totales?)")
            )
            continue
        try:
            registros.append(
                RegistroExcel(
                    fila=numero,
                    nit=parsear_nit(nit_raw),
                    nombre=str(nombre_raw or "").strip(),
                    valor=parsear_valor_excel(valor_raw),
                )
            )
        except ValueError as e:
            invalidas.append(FilaInvalida(fila=numero, motivo=str(e)))
    libro.close()
    return registros, invalidas


class ResultadosExcel:
    """Excel de salida: una fila por registro, guardado tras cada fila para no perder avance."""

    COLUMNAS = [
        "Fila Excel", "NIT", "Nombre", "Valor Excel",
        "Resultado", "Motivo", "Procesado",
    ]

    def __init__(self, directorio: Path):
        directorio.mkdir(exist_ok=True)
        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.ruta = directorio / f"resultado_{marca}.xlsx"
        self._libro = Workbook()
        self._hoja = self._libro.active
        self._hoja.title = "Resultado"
        self._hoja.append(self.COLUMNAS)
        self._libro.save(self.ruta)

    def agregar(self, registro: RegistroExcel, decision: Decision, motivo: str) -> None:
        self._hoja.append([
            registro.fila, registro.nit, registro.nombre, float(registro.valor),
            decision.value, motivo, datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ])
        self._libro.save(self.ruta)

    def agregar_invalida(self, invalida: FilaInvalida) -> None:
        self._hoja.append([
            invalida.fila, "", "", "", "fila inválida: revisar a mano", invalida.motivo,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ])
        self._libro.save(self.ruta)