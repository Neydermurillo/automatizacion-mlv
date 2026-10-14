from openpyxl import load_workbook

from src.config import cargar_config
from src.excel_io import leer_registros

config = cargar_config()

libro = load_workbook(config.archivo_excel, data_only=True)
hoja = libro[config.hoja_excel] if config.hoja_excel else libro.active
print(f"Hojas del libro: {libro.sheetnames}")
print(f"Hoja que se lee: {hoja.title} | filas según Excel: {hoja.max_row}")
libro.close()

registros, invalidas = leer_registros(config)
print(f"Registros válidos: {len(registros)} | filas inválidas: {len(invalidas)}")
print("--- primeros 3 ---")
for r in registros[:3]:
    print(r)
print("--- últimos 3 ---")
for r in registros[-3:]:
    print(r)
print("--- inválidas (hasta 10) ---")
for i in invalidas[:10]:
    print(i)